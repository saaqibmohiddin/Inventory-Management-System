import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook
import sqlite3


# ==========================================================
# DATABASE CONNECTION
# ==========================================================

def get_connection():

    return sqlite3.connect("inventory.db")


# ==========================================================
# GET REPORT DATA
# ==========================================================

def get_report_data():

    conn = get_connection()
    cursor = conn.cursor()

    # ------------------------------------------------------
    # TOTAL SALES
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COALESCE(SUM(total), 0)
        FROM sales
    """)

    total_sales = cursor.fetchone()[0]

    # ------------------------------------------------------
    # TOTAL PURCHASES
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COALESCE(SUM(total_cost), 0)
        FROM purchases
    """)

    total_purchases = cursor.fetchone()[0]

    # ------------------------------------------------------
    # TOTAL PROFIT
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COALESCE(
            SUM(
                total -
                (quantity * cost_price)
            ),
            0
        )
        FROM sales
    """)

    total_profit = cursor.fetchone()[0]

    # ------------------------------------------------------
    # CURRENT STOCK VALUE
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COALESCE(
            SUM(
                quantity * cost_price
            ),
            0
        )
        FROM products
    """)

    stock_value = cursor.fetchone()[0]

    # ------------------------------------------------------
    # TOTAL PRODUCTS
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COUNT(*)
        FROM products
    """)

    total_products = cursor.fetchone()[0]

    # ------------------------------------------------------
    # TOTAL STOCK
    # ------------------------------------------------------

    cursor.execute("""
        SELECT COALESCE(
            SUM(quantity),
            0
        )
        FROM products
    """)

    total_stock = cursor.fetchone()[0]

    # ------------------------------------------------------
    # LOW STOCK PRODUCTS
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            name,
            category,
            quantity
        FROM products
        WHERE quantity > 0
        AND quantity <= 5
        ORDER BY quantity ASC
    """)

    low_stock_products = cursor.fetchall()

    # ------------------------------------------------------
    # OUT OF STOCK PRODUCTS
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            name,
            category,
            quantity
        FROM products
        WHERE quantity = 0
        ORDER BY name
    """)

    out_of_stock_products = cursor.fetchall()

    # ------------------------------------------------------
    # TOP SELLING PRODUCTS
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            product_name,
            SUM(quantity) AS total_quantity,
            SUM(total) AS total_revenue
        FROM sales
        GROUP BY product_name
        ORDER BY total_quantity DESC
        LIMIT 10
    """)

    top_selling_products = cursor.fetchall()

    # ------------------------------------------------------
    # CATEGORY-WISE STOCK
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            category,
            COUNT(*) AS products,
            COALESCE(SUM(quantity), 0) AS stock
        FROM products
        GROUP BY category
        ORDER BY stock DESC
    """)

    category_stock = cursor.fetchall()

    # ------------------------------------------------------
    # RECENT SALES
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            product_name,
            quantity,
            total,
            sale_date
        FROM sales
        ORDER BY id DESC
        LIMIT 10
    """)

    recent_sales = cursor.fetchall()

    # ------------------------------------------------------
    # RECENT PURCHASES
    # ------------------------------------------------------

    cursor.execute("""
        SELECT
            id,
            product_name,
            quantity,
            total_cost,
            purchase_date
        FROM purchases
        ORDER BY id DESC
        LIMIT 10
    """)

    recent_purchases = cursor.fetchall()

    conn.close()

    return {
        "total_sales": total_sales,
        "total_purchases": total_purchases,
        "total_profit": total_profit,
        "stock_value": stock_value,
        "total_products": total_products,
        "total_stock": total_stock,
        "low_stock_products": low_stock_products,
        "out_of_stock_products": out_of_stock_products,
        "top_selling_products": top_selling_products,
        "category_stock": category_stock,
        "recent_sales": recent_sales,
        "recent_purchases": recent_purchases
    }


# ==========================================================
# EXPORT COMPLETE REPORT
# ==========================================================

def export_report():

    data = get_report_data()

    file_path = filedialog.asksaveasfilename(
        title="Save Business Report",
        defaultextension=".xlsx",
        filetypes=[
            ("Excel Files", "*.xlsx")
        ],
        initialfile="business_report.xlsx"
    )

    if not file_path:

        return

    try:

        workbook = Workbook()

        # ==================================================
        # SUMMARY SHEET
        # ==================================================

        summary = workbook.active

        summary.title = "Business Summary"

        summary.append([
            "INVENTORY MANAGEMENT SYSTEM"
        ])

        summary.append([])

        summary.append([
            "Metric",
            "Value"
        ])

        summary.append([
            "Total Products",
            data["total_products"]
        ])

        summary.append([
            "Total Stock",
            data["total_stock"]
        ])

        summary.append([
            "Total Sales",
            data["total_sales"]
        ])

        summary.append([
            "Total Purchases",
            data["total_purchases"]
        ])

        summary.append([
            "Total Profit",
            data["total_profit"]
        ])

        summary.append([
            "Current Stock Value",
            data["stock_value"]
        ])

        summary.column_dimensions["A"].width = 30

        summary.column_dimensions["B"].width = 25

        # ==================================================
        # TOP PRODUCTS SHEET
        # ==================================================

        top_sheet = workbook.create_sheet(
            "Top Selling"
        )

        top_sheet.append([
            "Product",
            "Quantity Sold",
            "Revenue"
        ])

        for product in data[
            "top_selling_products"
        ]:

            top_sheet.append([
                product[0],
                product[1],
                product[2]
            ])

        # ==================================================
        # LOW STOCK SHEET
        # ==================================================

        low_sheet = workbook.create_sheet(
            "Low Stock"
        )

        low_sheet.append([
            "ID",
            "Product",
            "Category",
            "Quantity"
        ])

        for product in data[
            "low_stock_products"
        ]:

            low_sheet.append([
                product[0],
                product[1],
                product[2],
                product[3]
            ])

        # ==================================================
        # OUT OF STOCK SHEET
        # ==================================================

        out_sheet = workbook.create_sheet(
            "Out of Stock"
        )

        out_sheet.append([
            "ID",
            "Product",
            "Category",
            "Quantity"
        ])

        for product in data[
            "out_of_stock_products"
        ]:

            out_sheet.append([
                product[0],
                product[1],
                product[2],
                product[3]
            ])

        # ==================================================
        # CATEGORY SHEET
        # ==================================================

        category_sheet = workbook.create_sheet(
            "Category Stock"
        )

        category_sheet.append([
            "Category",
            "Products",
            "Total Stock"
        ])

        for category in data[
            "category_stock"
        ]:

            category_sheet.append([
                category[0],
                category[1],
                category[2]
            ])

        # ==================================================
        # RECENT SALES SHEET
        # ==================================================

        sales_sheet = workbook.create_sheet(
            "Recent Sales"
        )

        sales_sheet.append([
            "Sale ID",
            "Product",
            "Quantity",
            "Total",
            "Date & Time"
        ])

        for sale in data[
            "recent_sales"
        ]:

            sales_sheet.append([
                sale[0],
                sale[1],
                sale[2],
                sale[3],
                sale[4]
            ])

        # ==================================================
        # RECENT PURCHASES SHEET
        # ==================================================

        purchase_sheet = workbook.create_sheet(
            "Recent Purchases"
        )

        purchase_sheet.append([
            "Purchase ID",
            "Product",
            "Quantity",
            "Total Cost",
            "Date & Time"
        ])

        for purchase in data[
            "recent_purchases"
        ]:

            purchase_sheet.append([
                purchase[0],
                purchase[1],
                purchase[2],
                purchase[3],
                purchase[4]
            ])

        workbook.save(file_path)

        messagebox.showinfo(
            "Export Successful",
            "Business report exported successfully!\n\n"
            f"Saved at:\n{file_path}"
        )

    except Exception as error:

        messagebox.showerror(
            "Export Error",
            f"Could not export report.\n\n{error}"
        )


# ==========================================================
# REPORTS WINDOW
# ==========================================================

def open_reports(parent):

    report_window = tk.Toplevel(parent)

    report_window.title(
        "Reports & Analytics"
    )

    report_window.geometry(
        "1150x850"
    )

    report_window.resizable(
        False,
        False
    )

    # ======================================================
    # TITLE
    # ======================================================

    tk.Label(
        report_window,
        text="📊 REPORTS & ANALYTICS",
        font=("Arial", 24, "bold")
    ).pack(
        pady=15
    )

    # ======================================================
    # DASHBOARD CARDS
    # ======================================================

    cards_frame = tk.Frame(
        report_window
    )

    cards_frame.pack(
        pady=5
    )

    total_sales_label = tk.Label(
        cards_frame,
        text="TOTAL SALES\n₹0.00",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    total_sales_label.grid(
        row=0,
        column=0,
        padx=5
    )

    total_purchases_label = tk.Label(
        cards_frame,
        text="TOTAL PURCHASES\n₹0.00",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    total_purchases_label.grid(
        row=0,
        column=1,
        padx=5
    )

    total_profit_label = tk.Label(
        cards_frame,
        text="TOTAL PROFIT\n₹0.00",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    total_profit_label.grid(
        row=0,
        column=2,
        padx=5
    )

    stock_value_label = tk.Label(
        cards_frame,
        text="STOCK VALUE\n₹0.00",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    stock_value_label.grid(
        row=0,
        column=3,
        padx=5
    )

    # ======================================================
    # SECOND ROW CARDS
    # ======================================================

    second_cards = tk.Frame(
        report_window
    )

    second_cards.pack(
        pady=5
    )

    total_products_label = tk.Label(
        second_cards,
        text="TOTAL PRODUCTS\n0",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    total_products_label.grid(
        row=0,
        column=0,
        padx=5
    )

    total_stock_label = tk.Label(
        second_cards,
        text="TOTAL STOCK\n0",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    total_stock_label.grid(
        row=0,
        column=1,
        padx=5
    )

    low_stock_label = tk.Label(
        second_cards,
        text="LOW STOCK\n0",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    low_stock_label.grid(
        row=0,
        column=2,
        padx=5
    )

    out_stock_label = tk.Label(
        second_cards,
        text="OUT OF STOCK\n0",
        font=("Arial", 11, "bold"),
        width=21,
        height=3,
        relief="ridge"
    )

    out_stock_label.grid(
        row=0,
        column=3,
        padx=5
    )

    # ======================================================
    # MAIN CONTENT
    # ======================================================

    content_frame = tk.Frame(
        report_window
    )

    content_frame.pack(
        fill="both",
        expand=True,
        padx=20,
        pady=10
    )

    # ======================================================
    # TOP SELLING PRODUCTS
    # ======================================================

    top_frame = tk.LabelFrame(
        content_frame,
        text="🏆 Top Selling Products",
        font=("Arial", 11, "bold")
    )

    top_frame.grid(
        row=0,
        column=0,
        padx=5,
        pady=5
    )

    top_columns = (
        "Product",
        "Quantity",
        "Revenue"
    )

    top_table = ttk.Treeview(
        top_frame,
        columns=top_columns,
        show="headings",
        height=8
    )

    top_table.heading(
        "Product",
        text="Product"
    )

    top_table.heading(
        "Quantity",
        text="Qty Sold"
    )

    top_table.heading(
        "Revenue",
        text="Revenue"
    )

    top_table.column(
        "Product",
        width=210
    )

    top_table.column(
        "Quantity",
        width=90,
        anchor="center"
    )

    top_table.column(
        "Revenue",
        width=120,
        anchor="center"
    )

    top_table.pack(
        padx=5,
        pady=5
    )

    # ======================================================
    # LOW STOCK
    # ======================================================

    low_frame = tk.LabelFrame(
        content_frame,
        text="⚠️ Low Stock Products",
        font=("Arial", 11, "bold")
    )

    low_frame.grid(
        row=0,
        column=1,
        padx=5,
        pady=5
    )

    low_columns = (
        "ID",
        "Product",
        "Category",
        "Stock"
    )

    low_table = ttk.Treeview(
        low_frame,
        columns=low_columns,
        show="headings",
        height=8
    )

    for column in low_columns:

        low_table.heading(
            column,
            text=column
        )

    low_table.column(
        "ID",
        width=50,
        anchor="center"
    )

    low_table.column(
        "Product",
        width=160
    )

    low_table.column(
        "Category",
        width=120
    )

    low_table.column(
        "Stock",
        width=70,
        anchor="center"
    )

    low_table.pack(
        padx=5,
        pady=5
    )

    # ======================================================
    # CATEGORY STOCK
    # ======================================================

    category_frame = tk.LabelFrame(
        content_frame,
        text="📂 Category-wise Stock",
        font=("Arial", 11, "bold")
    )

    category_frame.grid(
        row=1,
        column=0,
        padx=5,
        pady=5
    )

    category_columns = (
        "Category",
        "Products",
        "Stock"
    )

    category_table = ttk.Treeview(
        category_frame,
        columns=category_columns,
        show="headings",
        height=7
    )

    for column in category_columns:

        category_table.heading(
            column,
            text=column
        )

    category_table.column(
        "Category",
        width=200
    )

    category_table.column(
        "Products",
        width=100,
        anchor="center"
    )

    category_table.column(
        "Stock",
        width=100,
        anchor="center"
    )

    category_table.pack(
        padx=5,
        pady=5
    )

    # ======================================================
    # RECENT SALES
    # ======================================================

    recent_frame = tk.LabelFrame(
        content_frame,
        text="🕐 Recent Sales",
        font=("Arial", 11, "bold")
    )

    recent_frame.grid(
        row=1,
        column=1,
        padx=5,
        pady=5
    )

    recent_columns = (
        "ID",
        "Product",
        "Qty",
        "Total",
        "Date"
    )

    recent_table = ttk.Treeview(
        recent_frame,
        columns=recent_columns,
        show="headings",
        height=7
    )

    for column in recent_columns:

        recent_table.heading(
            column,
            text=column
        )

    recent_table.column(
        "ID",
        width=50,
        anchor="center"
    )

    recent_table.column(
        "Product",
        width=150
    )

    recent_table.column(
        "Qty",
        width=60,
        anchor="center"
    )

    recent_table.column(
        "Total",
        width=100,
        anchor="center"
    )

    recent_table.column(
        "Date",
        width=150,
        anchor="center"
    )

    recent_table.pack(
        padx=5,
        pady=5
    )

    # ======================================================
    # REFRESH REPORT
    # ======================================================

    def refresh_report():

        data = get_report_data()

        # Dashboard
        total_sales_label.config(
            text=f"TOTAL SALES\n₹{data['total_sales']:,.2f}"
        )

        total_purchases_label.config(
            text=f"TOTAL PURCHASES\n₹{data['total_purchases']:,.2f}"
        )

        total_profit_label.config(
            text=f"TOTAL PROFIT\n₹{data['total_profit']:,.2f}"
        )

        stock_value_label.config(
            text=f"STOCK VALUE\n₹{data['stock_value']:,.2f}"
        )

        total_products_label.config(
            text=f"TOTAL PRODUCTS\n{data['total_products']}"
        )

        total_stock_label.config(
            text=f"TOTAL STOCK\n{data['total_stock']}"
        )

        low_stock_label.config(
            text=f"LOW STOCK\n{len(data['low_stock_products'])}"
        )

        out_stock_label.config(
            text=f"OUT OF STOCK\n{len(data['out_of_stock_products'])}"
        )

        # --------------------------------------------------
        # Clear top selling
        # --------------------------------------------------

        for item in top_table.get_children():

            top_table.delete(item)

        # Add top products

        for product in data[
            "top_selling_products"
        ]:

            top_table.insert(
                "",
                tk.END,
                values=(
                    product[0],
                    product[1],
                    f"₹{float(product[2]):,.2f}"
                )
            )

        # --------------------------------------------------
        # Clear low stock
        # --------------------------------------------------

        for item in low_table.get_children():

            low_table.delete(item)

        for product in data[
            "low_stock_products"
        ]:

            low_table.insert(
                "",
                tk.END,
                values=product
            )

        # --------------------------------------------------
        # Clear category
        # --------------------------------------------------

        for item in category_table.get_children():

            category_table.delete(item)

        for category in data[
            "category_stock"
        ]:

            category_table.insert(
                "",
                tk.END,
                values=category
            )

        # --------------------------------------------------
        # Clear recent sales
        # --------------------------------------------------

        for item in recent_table.get_children():

            recent_table.delete(item)

        for sale in data[
            "recent_sales"
        ]:

            recent_table.insert(
                "",
                tk.END,
                values=(
                    sale[0],
                    sale[1],
                    sale[2],
                    f"₹{float(sale[3]):,.2f}",
                    sale[4]
                )
            )

    # ======================================================
    # BUTTONS
    # ======================================================

    button_frame = tk.Frame(
        report_window
    )

    button_frame.pack(
        pady=10
    )

    tk.Button(
        button_frame,
        text="🔄 Refresh Report",
        width=18,
        command=refresh_report
    ).grid(
        row=0,
        column=0,
        padx=5
    )

    tk.Button(
        button_frame,
        text="📥 Export Full Report",
        width=22,
        command=export_report
    ).grid(
        row=0,
        column=1,
        padx=5
    )

    tk.Button(
        button_frame,
        text="CLOSE",
        width=15,
        command=report_window.destroy
    ).grid(
        row=0,
        column=2,
        padx=5
    )

    # First load
    refresh_report()